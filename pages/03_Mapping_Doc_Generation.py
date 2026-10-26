"""
Streamlit in Snowflake page: Source-to-Target Mapping Generator

Place this file in your app's pages/ folder, e.g.
    pages/3_Source_Target_Mapping.py
It reads the same config.json the rest of the app uses.
"""

import io
import json
import re
from datetime import datetime
from pathlib import Path

import pandas as pd
import streamlit as st
from snowflake.snowpark.context import get_active_session

# ============================================================
# CONFIG
# ============================================================
def load_config() -> dict:
    """Find config.json in config/ (app root), the app root, or next to this page."""
    here = Path(__file__).resolve().parent
    candidates = (
        here.parent / "config" / "config.json",   # <app root>/config/config.json
        here.parent / "config.json",              # <app root>/config.json
        here / "config" / "config.json",
        here / "config.json",
    )
    for candidate in candidates:
        if candidate.exists():
            with open(candidate, "r") as f:
                return json.load(f)
    raise FileNotFoundError(
        "config.json not found. Looked in: " + ", ".join(str(c) for c in candidates)
    )


CONFIG = load_config()
APP_CFG = CONFIG.get("application", {})
AI_CFG = CONFIG.get("ai", {})
MAP_CFG = CONFIG.get("mapping", {})

SOURCE_DATABASE = MAP_CFG.get("source_database", "HDMP_SILVER")
TARGET_DATABASE = MAP_CFG.get("target_database", "HDMP_GOLD")
MODEL = AI_CFG.get("model", "claude-sonnet-5")
MAX_RETRIES = int(AI_CFG.get("max_repair_attempts", 1))
MAX_TOKENS = int(AI_CFG.get("max_tokens", 15000))
AVAILABLE_SCHEMAS = list(CONFIG.get("sources", {}).keys()) or ["JSON", "XML"]

OUTPUT_COLUMNS = [
    "SOURCE_DATABASE", "SOURCE_SCHEMA", "SOURCE_TABLE", "SOURCE_COLUMN",
    "SOURCE_DATATYPE", "SOURCE_NULLABLE",
    "TARGET_DATABASE", "TARGET_SCHEMA", "TARGET_TABLE", "TARGET_COLUMN",
    "TARGET_DATATYPE", "TARGET_NULLABLE",
    "MAPPING_RULE", "NOTES", "TABLE_DESCRIPTION", "BUSINESS_RULE",
    "CONFIDENCE", "COMMENTS",
]

session = get_active_session()


# ============================================================
# METADATA
# ============================================================
@st.cache_data(ttl=300, show_spinner=False)
def fetch_schema_metadata(database: str, schema: str) -> dict:
    """{table: {column: {"type": ..., "nullable": ...}}} for base tables in a schema."""
    query = f"""
        SELECT c.TABLE_NAME, c.COLUMN_NAME, c.DATA_TYPE, c.IS_NULLABLE
        FROM {database}.INFORMATION_SCHEMA.COLUMNS c
        JOIN {database}.INFORMATION_SCHEMA.TABLES t
          ON c.TABLE_CATALOG = t.TABLE_CATALOG
         AND c.TABLE_SCHEMA  = t.TABLE_SCHEMA
         AND c.TABLE_NAME    = t.TABLE_NAME
        WHERE c.TABLE_SCHEMA = '{schema}'
          AND t.TABLE_TYPE = 'BASE TABLE'
        ORDER BY c.TABLE_NAME, c.ORDINAL_POSITION
    """
    rows = session.sql(query).collect()
    out: dict = {}
    for r in rows:
        out.setdefault(r["TABLE_NAME"], {})[r["COLUMN_NAME"]] = {
            "type": r["DATA_TYPE"],
            "nullable": r["IS_NULLABLE"],
        }
    return out


# ============================================================
# CORTEX
# ============================================================
def build_prompt(schema: str, table: str, columns: dict) -> str:
    return f"""
You are a senior data mapping engineer.

Create a source-to-target mapping for the source table provided below.

SOURCE DATABASE: {SOURCE_DATABASE}
SOURCE SCHEMA: {schema}
SOURCE TABLE: {table}

SOURCE TABLE COLUMNS (name -> data type and nullability):
{json.dumps(columns, indent=2)}

TARGET DATABASE: {TARGET_DATABASE}
TARGET SCHEMA: {schema}

Rules for the target:
- Use the target database and the same schema as the source.
- Use the source table as the basis for the target table.
- The target data model should be a proper presentation layer.
- Determine appropriate target columns and transformations from the source metadata.
- Produce one JSON object per target column.

Return ONLY a JSON array. No markdown fences, no commentary.

Each JSON object must contain exactly these fields:
{", ".join(OUTPUT_COLUMNS)}

MAPPING_RULE must be one of: DIRECT, TRANSFORMATION, DATATYPE_CONVERSION, MULTIPLE_SOURCE, UNKNOWN
CONFIDENCE must be one of: HIGH, MEDIUM, LOW
"""


def call_cortex(prompt: str) -> str | None:
    query = """
        SELECT AI_COMPLETE(
            model => ?,
            prompt => ?,
            model_parameters => {'max_tokens': %d}
        ) AS RESPONSE
    """ % MAX_TOKENS
    result = session.sql(query, params=[MODEL, prompt]).collect()
    if not result or not result[0]["RESPONSE"]:
        return None
    return result[0]["RESPONSE"]


def extract_text(response, _depth: int = 0) -> str:
    """Unwrap the AI_COMPLETE envelope / double-encoded JSON string if present."""
    if _depth > 3:
        return response
    try:
        obj = json.loads(response)
    except (json.JSONDecodeError, TypeError):
        return response
    if isinstance(obj, str):
        return extract_text(obj, _depth + 1)
    if isinstance(obj, dict) and obj.get("choices"):
        choice = obj["choices"][0]
        text = choice.get("messages") or (choice.get("message") or {}).get("content")
        if isinstance(text, str) and text:
            return extract_text(text, _depth + 1)
    return response


def parse_mapping_json(text: str):
    """
    Returns (rows, partial).
    1) Try every '[' as the start of a complete JSON array of objects.
    2) If none decodes (e.g. output truncated by max_tokens), salvage every
       complete {...} object after the first '[' and flag the result as partial.
    """
    text = re.sub(r"```(?:json)?", "", text).strip()
    dec = json.JSONDecoder()

    for m in re.finditer(r"\[", text):
        try:
            obj, _ = dec.raw_decode(text, m.start())
        except json.JSONDecodeError:
            continue
        if isinstance(obj, list) and obj and all(isinstance(o, dict) for o in obj):
            return obj, False

    m = re.search(r"\[\s*\{", text)
    if m:
        pos = m.end() - 1
        rows = []
        while pos < len(text):
            while pos < len(text) and text[pos] in " \t\r\n,":
                pos += 1
            if pos >= len(text) or text[pos] != "{":
                break
            try:
                obj, pos = dec.raw_decode(text, pos)
            except json.JSONDecodeError:
                break
            if isinstance(obj, dict):
                rows.append(obj)
        if rows:
            return rows, True

    raise ValueError("Could not find a valid JSON array in the response")


STRICT_SUFFIX = (
    "\n\nIMPORTANT: Your previous reply was not valid JSON. Return ONLY the raw "
    "JSON array, starting with [ and ending with ]. No commentary, no markdown."
)


def generate_for_table(schema: str, table: str, columns: dict):
    """Returns (rows, error, partial, raw_response)."""
    base_prompt = build_prompt(schema, table, columns)
    last_err, last_raw = None, ""
    best_partial: list[dict] = []

    for attempt in range(MAX_RETRIES + 1):
        prompt = base_prompt if attempt == 0 else base_prompt + STRICT_SUFFIX
        try:
            raw = call_cortex(prompt)
            if not raw:
                last_err = "Empty response from Cortex"
                continue
            last_raw = raw
            rows, partial = parse_mapping_json(extract_text(raw))
            if not partial:
                return rows, None, False, last_raw
            if len(rows) > len(best_partial):
                best_partial = rows
            last_err = "Response truncated (hit max_tokens)"
        except Exception as e:  # noqa: BLE001
            last_err = str(e)

    if best_partial:
        return best_partial, None, True, last_raw
    return [], last_err, False, last_raw


def normalize_rows(rows: list[dict], schema: str, table: str) -> list[dict]:
    """Force deterministic fields and make sure every output column exists."""
    clean = []
    for r in rows:
        row = {c: r.get(c, "") for c in OUTPUT_COLUMNS}
        row["SOURCE_DATABASE"] = SOURCE_DATABASE
        row["SOURCE_SCHEMA"] = schema
        row["SOURCE_TABLE"] = table
        row["TARGET_DATABASE"] = TARGET_DATABASE
        row["TARGET_SCHEMA"] = schema
        for k, v in row.items():
            if isinstance(v, (list, dict)):
                row[k] = json.dumps(v)
            elif v is None:
                row[k] = ""
        clean.append(row)
    return clean


# ============================================================
# EXCEL
# ============================================================
def build_excel(df: pd.DataFrame, summary: pd.DataFrame, failed: pd.DataFrame) -> bytes:
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Mapping", index=False)
        summary.to_excel(writer, sheet_name="Run Summary", index=False)
        if not failed.empty:
            failed.to_excel(writer, sheet_name="Failed or Partial", index=False)

        header_fill = PatternFill("solid", fgColor="1F4E78")
        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
                cell.alignment = Alignment(vertical="center")
            for idx, col in enumerate(ws.columns, start=1):
                longest = max((len(str(c.value)) for c in col if c.value is not None), default=10)
                ws.column_dimensions[get_column_letter(idx)].width = min(max(longest + 2, 12), 60)
            ws.auto_filter.ref = ws.dimensions
    return buf.getvalue()


# ============================================================
# UI
# ============================================================
st.title("Source-to-Target Mapping")
st.caption(
    f"Generate mappings from **{SOURCE_DATABASE}** to **{TARGET_DATABASE}** "
    f"using Snowflake Cortex ({MODEL})."
)

schemas = st.multiselect("Source schemas", AVAILABLE_SCHEMAS, default=[s for s in ("JSON", "XML") if s in AVAILABLE_SCHEMAS])

metadata: dict = {}
for s in schemas:
    try:
        metadata[s] = fetch_schema_metadata(SOURCE_DATABASE, s)
    except Exception as e:  # noqa: BLE001
        st.error(f"Could not read metadata for {SOURCE_DATABASE}.{s}: {e}")

# Table selection (defaults to all tables in the chosen schemas)
all_tables = [f"{s}.{t}" for s, tables in metadata.items() for t in tables]
selected = st.multiselect("Tables to map", all_tables, default=all_tables)

if schemas:
    st.caption(f"{len(selected)} of {len(all_tables)} tables selected")

run = st.button("Generate mapping", type="primary", disabled=not selected)

if run:
    all_rows: list[dict] = []
    summary_rows: list[dict] = []
    failed_rows: list[dict] = []

    progress = st.progress(0.0)
    status = st.empty()

    for i, full_name in enumerate(selected, start=1):
        schema, table = full_name.split(".", 1)
        columns = metadata[schema][table]
        status.write(f"[{i}/{len(selected)}] Processing **{full_name}** ({len(columns)} columns)...")

        rows, err, partial, raw = generate_for_table(schema, table, columns)
        if err:
            failed_rows.append({
                "SOURCE_SCHEMA": schema, "SOURCE_TABLE": table, "ERROR": err,
                "RAW_RESPONSE": (raw or "")[:5000],
            })
            summary_rows.append({"SOURCE_SCHEMA": schema, "SOURCE_TABLE": table,
                                 "SOURCE_COLUMNS": len(columns), "MAPPING_ROWS": 0, "STATUS": "FAILED"})
        else:
            clean = normalize_rows(rows, schema, table)
            all_rows.extend(clean)
            if partial:
                failed_rows.append({
                    "SOURCE_SCHEMA": schema, "SOURCE_TABLE": table,
                    "ERROR": "PARTIAL: response truncated, some columns may be missing",
                    "RAW_RESPONSE": (raw or "")[:5000],
                })
            summary_rows.append({"SOURCE_SCHEMA": schema, "SOURCE_TABLE": table,
                                 "SOURCE_COLUMNS": len(columns), "MAPPING_ROWS": len(clean),
                                 "STATUS": "PARTIAL" if partial else "SUCCESS"})
        progress.progress(i / len(selected))

    status.empty()
    st.session_state["stm_df"] = pd.DataFrame(all_rows, columns=OUTPUT_COLUMNS)
    st.session_state["stm_summary"] = pd.DataFrame(summary_rows)
    st.session_state["stm_failed"] = pd.DataFrame(failed_rows)
    st.session_state["stm_generated_at"] = datetime.now().strftime("%Y%m%d_%H%M%S")

# ------------------------------------------------------------
# Results (kept in session_state so the download doesn't wipe them)
# ------------------------------------------------------------
if "stm_df" in st.session_state:
    df = st.session_state["stm_df"]
    summary = st.session_state["stm_summary"]
    failed = st.session_state["stm_failed"]

    ok = int((summary["STATUS"] == "SUCCESS").sum())
    bad = int((summary["STATUS"] == "FAILED").sum())

    partial_n = int((summary["STATUS"] == "PARTIAL").sum())
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Succeeded", ok)
    c2.metric("Partial", partial_n)
    c3.metric("Failed", bad)
    c4.metric("Mapping rows", len(df))

    if not df.empty:
        st.subheader("Mapping preview")
        st.dataframe(df, use_container_width=True, height=450)

        st.download_button(
            label="Download as Excel",
            data=build_excel(df, summary, failed),
            file_name=f"source_target_mapping_{st.session_state['stm_generated_at']}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
        )

    if not failed.empty:
        with st.expander(f"Failed / partial tables ({len(failed)})"):
            st.dataframe(failed.drop(columns=["RAW_RESPONSE"], errors="ignore"), use_container_width=True)
            for _, r in failed.iterrows():
                if r.get("RAW_RESPONSE"):
                    st.markdown(f"**Raw Cortex response: {r['SOURCE_SCHEMA']}.{r['SOURCE_TABLE']}**")
                    st.code(r["RAW_RESPONSE"], language="text")