import json
import re
from io import BytesIO

import pandas as pd
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


MAPPING_COLUMNS = [
    "SOURCE_DATABASE",
    "SOURCE_SCHEMA",
    "SOURCE_TABLE",
    "SOURCE_COLUMN",
    "SOURCE_DATATYPE",
    "SOURCE_NULLABLE",
    "TARGET_DATABASE",
    "TARGET_SCHEMA",
    "TARGET_TABLE",
    "TARGET_COLUMN",
    "TARGET_DATATYPE",
    "TARGET_NULLABLE",
    "MAPPING_RULE",
    "NOTES",
    "TABLE_DESCRIPTION",
    "BUSINESS_RULE",
    "CONFIDENCE",
    "COMMENTS",
]


def quote_identifier(identifier: str) -> str:
    return '"' + identifier.replace('"', '""') + '"'


def fetch_schema_metadata(session, source_database: str, schema: str) -> dict:
    """Fetch base-table columns, data types and nullability."""
    database = quote_identifier(source_database)
    schema_literal = schema.replace("'", "''")

    query = f"""
        SELECT
            c.TABLE_NAME,
            c.COLUMN_NAME,
            c.DATA_TYPE,
            c.IS_NULLABLE
        FROM {database}.INFORMATION_SCHEMA.COLUMNS c
        JOIN {database}.INFORMATION_SCHEMA.TABLES t
          ON c.TABLE_CATALOG = t.TABLE_CATALOG
         AND c.TABLE_SCHEMA = t.TABLE_SCHEMA
         AND c.TABLE_NAME = t.TABLE_NAME
        WHERE c.TABLE_SCHEMA = '{schema_literal}'
          AND t.TABLE_TYPE = 'BASE TABLE'
        ORDER BY
            c.TABLE_NAME,
            c.ORDINAL_POSITION
    """

    rows = session.sql(query).collect()

    schema_dict = {}

    for row in rows:
        table_name = row["TABLE_NAME"]
        column_name = row["COLUMN_NAME"]

        schema_dict.setdefault(table_name, {})
        schema_dict[table_name][column_name] = {
            "DATA_TYPE": row["DATA_TYPE"],
            "IS_NULLABLE": row["IS_NULLABLE"],
        }

    return schema_dict


def get_all_metadata(
    session,
    source_database: str,
    source_schemas: list[str],
) -> dict:
    """Fetch metadata for all selected schemas."""
    return {
        schema: fetch_schema_metadata(
            session,
            source_database,
            schema,
        )
        for schema in source_schemas
    }


def build_table_metadata_for_prompt(
    table_name: str,
    columns: dict,
) -> dict:
    return {
        table_name: columns,
    }


def call_cortex(
    session,
    table_metadata: dict,
    source_database: str,
    source_schema: str,
    source_table: str,
    target_database: str,
    model: str,
    max_tokens: int,
) :
    """Call Snowflake Cortex once for one source table."""
    prompt = f"""
You are a senior data mapping engineer.

Create a source-to-target mapping for the source table provided below.

SOURCE DATABASE:
{source_database}

SOURCE SCHEMA:
{source_schema}

SOURCE TABLE:
{source_table}

SOURCE TABLE COLUMNS, DATA TYPES AND NULLABILITY:
{json.dumps(table_metadata, indent=2)}

TARGET DATABASE:
{target_database}

TARGET SCHEMA:
{source_schema}

Create a source-to-target mapping.

For the target:
- Use the target database mentioned above.
- Use the same schema as the source.
- Use the source table as the basis for the target table.
- The target data model should be a proper presentation layer.
- Determine appropriate target columns and transformations based on the source metadata.
- Do not invent source columns that are not present in the source metadata.
- Preserve source column names when the mapping is DIRECT.
- If a transformation is required, describe it clearly in MAPPING_RULE, NOTES, and BUSINESS_RULE.

Return ONLY a JSON array.

Each JSON object must contain exactly these fields:

SOURCE_DATABASE
SOURCE_SCHEMA
SOURCE_TABLE
SOURCE_COLUMN
SOURCE_DATATYPE
SOURCE_NULLABLE
TARGET_DATABASE
TARGET_SCHEMA
TARGET_TABLE
TARGET_COLUMN
TARGET_DATATYPE
TARGET_NULLABLE
MAPPING_RULE
NOTES
TABLE_DESCRIPTION
BUSINESS_RULE
CONFIDENCE
COMMENTS

MAPPING_RULE can be:
DIRECT
TRANSFORMATION
DATATYPE_CONVERSION
MULTIPLE_SOURCE
UNKNOWN

CONFIDENCE can be:
HIGH
MEDIUM
LOW
"""

    query = """
        SELECT AI_COMPLETE(
            model => ?,
            prompt => ?,
            model_parameters => {
                'max_tokens': ?
            }
        ) AS RESPONSE
    """

    result = session.sql(
        query,
        params=[model, prompt, max_tokens],
    ).collect()

    if not result:
        raise RuntimeError("No response returned from Cortex.")

    response = result[0]["RESPONSE"]

    if not response:
        raise RuntimeError("Cortex returned an empty response.")

    return response


def _extract_model_text(response):
    """Extract the actual model message from common AI_COMPLETE response shapes."""
    # Structured output may be returned as a one-item list containing
    # {"raw_message": ...}. Unwrap that before treating it as mapping rows.
    if isinstance(response, list):
        if len(response) == 1 and isinstance(response[0], dict):
            if "raw_message" in response[0]:
                return _extract_model_text(response[0]["raw_message"])
        return response

    # Snowpark can return a Python dict/object for JSON-like responses.
    if isinstance(response, dict):
        # AI_COMPLETE/COMPLETE detailed response shape
        choices = response.get("choices")
        if isinstance(choices, list) and choices:
            first_choice = choices[0]
            if isinstance(first_choice, dict) and "messages" in first_choice:
                return _extract_model_text(first_choice["messages"])

        # Structured-output shape
        if "structured_output" in response:
            return _extract_model_text(response["structured_output"])

        if "raw_message" in response:
            return _extract_model_text(response["raw_message"])

        # Some responses may already contain the mapping under a named key.
        for key in ("mappings", "mapping", "data", "response"):
            if key in response:
                return _extract_model_text(response[key])

        return response

    # Snowpark Row / VARIANT values may expose JSON through to_dict().
    if hasattr(response, "to_dict"):
        try:
            return _extract_model_text(response.to_dict())
        except Exception:
            pass

    return response


def _clean_json_text(text: str) -> str:
    """Remove common markdown wrappers around JSON."""
    cleaned = text.strip()

    # Remove fenced code blocks such as ```json ... ```.
    cleaned = re.sub(
        r"^\s*```(?:json)?\s*",
        "",
        cleaned,
        flags=re.IGNORECASE,
    )
    cleaned = re.sub(r"\s*```\s*$", "", cleaned)

    return cleaned.strip()


def parse_cortex_response(response) -> list[dict]:
    """Parse Cortex output even when AI_COMPLETE wraps the model message."""
    value = _extract_model_text(response)

    # Already a Python list.
    if isinstance(value, list):
        parsed = value
    else:
        # A JSON object/string can contain the actual model response.
        if isinstance(value, dict):
            if isinstance(value.get("mappings"), list):
                parsed = value["mappings"]
            else:
                # If this is a single mapping object, accept it as one row.
                if any(column in value for column in MAPPING_COLUMNS):
                    parsed = [value]
                else:
                    raise ValueError(
                        "Cortex returned a JSON object, but no mapping array "
                        "was found."
                    )
        else:
            cleaned = _clean_json_text(str(value))

            # Models sometimes return a JSON string containing another JSON
            # string/array. Decode repeatedly until the useful payload appears.
            parsed = None
            current = cleaned

            for _ in range(3):
                try:
                    candidate = json.loads(current)
                except json.JSONDecodeError:
                    candidate = None

                if isinstance(candidate, list):
                    parsed = candidate
                    break

                if isinstance(candidate, dict):
                    extracted = _extract_model_text(candidate)
                    if isinstance(extracted, list):
                        parsed = extracted
                        break
                    if isinstance(extracted, dict):
                        if isinstance(extracted.get("mappings"), list):
                            parsed = extracted["mappings"]
                            break
                        if any(column in extracted for column in MAPPING_COLUMNS):
                            parsed = [extracted]
                            break
                    if extracted is not candidate:
                        current = _clean_json_text(str(extracted))
                        continue

                if isinstance(candidate, str):
                    current = _clean_json_text(candidate)
                    continue

                break

            # Last-resort extraction for a response containing explanatory
            # text before/after the JSON array.
            if parsed is None:
                start = cleaned.find("[")
                end = cleaned.rfind("]")

                if start != -1 and end > start:
                    try:
                        candidate = json.loads(cleaned[start:end + 1])
                    except json.JSONDecodeError as exc:
                        raise ValueError(
                            "Cortex returned text containing an invalid JSON "
                            f"array: {exc}"
                        ) from exc

                    if isinstance(candidate, list):
                        parsed = candidate

            if parsed is None:
                preview = cleaned[:500].replace("\n", " ")
                raise ValueError(
                    "Cortex response could not be parsed as a mapping array. "
                    f"Response preview: {preview}"
                )

    rows = []

    for item in parsed:
        if not isinstance(item, dict):
            continue

        rows.append(
            {
                column: item.get(column, "")
                for column in MAPPING_COLUMNS
            }
        )

    if not rows:
        raise ValueError("Cortex returned an empty mapping array.")

    return rows

def generate_mapping(
    session,
    source_database: str,
    target_database: str,
    selected_tables: list[tuple[str, str]],
    metadata: dict,
    model: str,
    max_tokens: int,
    progress_callback=None,
    status_callback=None,
):
    """
    Generate mapping table-by-table.

    progress_callback(value) receives a float from 0 to 1.
    status_callback(message, status) receives status='info'|'success'|'error'.
    """
    mapping_rows = []
    run_details = []

    total = len(selected_tables)

    for index, (schema, table_name) in enumerate(
        selected_tables,
        start=1,
    ):
        if status_callback:
            status_callback(
                f"Processing {schema}.{table_name} "
                f"({index}/{total}) — calling Cortex...",
                "info",
            )

        try:
            table_columns = metadata[schema][table_name]

            table_metadata = build_table_metadata_for_prompt(
                table_name,
                table_columns,
            )

            response = call_cortex(
                session=session,
                table_metadata=table_metadata,
                source_database=source_database,
                source_schema=schema,
                source_table=table_name,
                target_database=target_database,
                model=model,
                max_tokens=max_tokens,
            )

            rows = parse_cortex_response(response)

            # Keep the workbook context controlled by the application.
            for row in rows:
                row["SOURCE_DATABASE"] = source_database
                row["SOURCE_SCHEMA"] = schema
                row["SOURCE_TABLE"] = table_name
                row["TARGET_DATABASE"] = target_database
                row["TARGET_SCHEMA"] = schema

            mapping_rows.extend(rows)

            run_details.append(
                {
                    "SOURCE_SCHEMA": schema,
                    "SOURCE_TABLE": table_name,
                    "STATUS": "SUCCESS",
                    "MAPPING_ROWS": len(rows),
                    "ERROR": "",
                }
            )

            if status_callback:
                status_callback(
                    f"Completed {schema}.{table_name} — "
                    f"{len(rows)} mapping rows generated.",
                    "success",
                )

        except Exception as exc:
            run_details.append(
                {
                    "SOURCE_SCHEMA": schema,
                    "SOURCE_TABLE": table_name,
                    "STATUS": "FAILED",
                    "MAPPING_ROWS": 0,
                    "ERROR": str(exc),
                }
            )

            if status_callback:
                status_callback(
                    f"Failed {schema}.{table_name} — {str(exc)}",
                    "error",
                )

        if progress_callback:
            progress_callback(index / total)

    return mapping_rows, run_details


def create_excel(
    mapping_rows: list[dict],
    run_details: list[dict],
) -> bytes:
    """Create the downloadable Excel mapping workbook."""
    output = BytesIO()

    mapping_df = pd.DataFrame(
        mapping_rows,
        columns=MAPPING_COLUMNS,
    )

    summary_df = pd.DataFrame(run_details)

    run_summary_df = pd.DataFrame(
        [
            {
                "TOTAL_TABLES": len(run_details),
                "SUCCESSFUL_TABLES": sum(
                    r["STATUS"] == "SUCCESS"
                    for r in run_details
                ),
                "FAILED_TABLES": sum(
                    r["STATUS"] == "FAILED"
                    for r in run_details
                ),
                "TOTAL_MAPPING_ROWS": len(mapping_rows),
            }
        ]
    )

    with pd.ExcelWriter(
        output,
        engine="openpyxl",
    ) as writer:
        mapping_df.to_excel(
            writer,
            sheet_name="Mapping",
            index=False,
        )

        summary_df.to_excel(
            writer,
            sheet_name="Table_Summary",
            index=False,
        )

        run_summary_df.to_excel(
            writer,
            sheet_name="Run_Summary",
            index=False,
        )

        workbook = writer.book

        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions

            for cell in worksheet[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )

            for column_cells in worksheet.columns:
                max_length = 0
                column_letter = get_column_letter(
                    column_cells[0].column
                )

                for cell in column_cells:
                    value = (
                        ""
                        if cell.value is None
                        else str(cell.value)
                    )
                    max_length = max(max_length, len(value))

                worksheet.column_dimensions[
                    column_letter
                ].width = min(
                    max(max_length + 2, 12),
                    45,
                )

        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    cell.alignment = Alignment(
                        vertical="top",
                        wrap_text=True,
                    )

    return output.getvalue()
